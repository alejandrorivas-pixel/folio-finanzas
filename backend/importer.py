"""Read-only adapter for the user's Balance workbook. Preserve every valid posting."""
import hashlib
import io
import json
import warnings
from collections import Counter
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile

import openpyxl

from . import ledger

SHEETS = {
    "Movimientos Efectivo": ("efectivo", "Balance Efectivo", "C5"),
    "Movimientos Banco": ("banco", "Balance Banco", "C5"),
    "Movimientos Tarjeta BBVA": ("tarjeta_bbva", "Balance Tarjeta BBVA", "C9"),
}


def refresh_preview(payload):
    """Recompute identity and counters after an explicit proposal correction."""
    occurrences = Counter()
    keys = ledger.existing_keys()
    order = {name: i for i, name in enumerate(SHEETS)}
    payload["movements"].sort(key=lambda r: (order.get(r.get("source_sheet"), 99), r.get("source_row", 0)))
    for r in payload["movements"]:
        canonical = {k: r.get(k) for k in ["date", "account", "kind", "description", "amount_minor", "reference", "property"]}
        base = hashlib.sha256(json.dumps(canonical, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        occurrences[base] += 1
        r["dedup_key"] = f"{base}:{occurrences[base]}"
        r["duplicate"] = r["dedup_key"] in keys
    new = [r for r in payload["movements"] if not r["duplicate"]]
    payload["stats"] = {**ledger.totals(new), "total": len(payload["movements"]), "new": len(new),
        "duplicates": len(payload["movements"]) - len(new), "errors": len(payload["errors"]),
        "accounts": dict(Counter(r["account"] for r in new)),
        "review": sum(r["treatment"] == "por_revisar" for r in new)}
    return payload


def parse_workbook(content, filename):
    if len(content) > 15 * 1024 * 1024:
        raise ValueError("El archivo excede 15 MB")
    try:
        with ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise ValueError("El contenido descomprimido excede el límite")
    except BadZipFile:
        raise ValueError("El archivo no es un Excel XLSX válido")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        w = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if not set(SHEETS).intersection(w.sheetnames):
            raise ValueError("No se encontraron las hojas de movimientos del Balance. Usa tu archivo original .xlsx.")
        sha = hashlib.sha256(content).hexdigest()
        movements, errors, notes, openings, categories = [], [], [], [], set()
        for name, (account, balance_sheet, cell) in SHEETS.items():
            if name not in w.sheetnames:
                notes.append(f"No se encontró {name}; no se importará esa cuenta.")
                continue
            sheet = w[name]
            if sheet.max_row > 50000:
                raise ValueError("La hoja excede 50,000 filas")
            header = [c.value for c in next(sheet.iter_rows(min_row=3, max_row=3, max_col=9))]
            if header[1:5] != ["Fecha", "Tipo", "Categoría", "Concepto / Descripción"]:
                raise ValueError(f"Los encabezados de {name} cambiaron; revisa el formato.")
            if balance_sheet in w.sheetnames and w[balance_sheet][cell].value is not None:
                opening = ledger.money(w[balance_sheet][cell].value)
                openings.append({"account": account, "amount_minor": -opening if account == "tarjeta_bbva" else opening,
                    "source": f"{filename} · {balance_sheet}!{cell}"})
            else:
                notes.append(f"Falta el saldo inicial de {account}; su saldo quedará sin determinar.")
            for idx, row in enumerate(sheet.iter_rows(min_row=4, max_col=17, values_only=True), 4):
                _, dt, kind, cat, desc, ref, prop, incoming, outgoing, *_ = row
                if dt is None and kind is None and ledger.normalize(cat) in {"total", "totales", "subtotal"}:
                    continue
                if not any(v is not None for v in (dt, kind, desc, incoming, outgoing)):
                    continue
                if kind is None and incoming is None and outgoing is None:
                    continue  # Opening balance note, not a transaction.
                try:
                    if not isinstance(dt, (date, datetime)):
                        raise ValueError("Fecha vacía o inválida")
                    kind = ledger.normalize(kind)
                    card = account == "tarjeta_bbva"
                    allowed = {"cargo", "pago", "interes", "comision", "abono", "devolucion"} if card else {"ingreso", "egreso"}
                    if kind not in allowed:
                        raise ValueError("Tipo de movimiento desconocido")
                    a, b = ledger.money(incoming or 0), ledger.money(outgoing or 0)
                    if a < 0 or b < 0 or (a and b) or not (a or b):
                        raise ValueError("Se requiere exactamente un importe positivo en entrada/cargo o salida/abono")
                    negative = kind in {"egreso", "cargo", "interes", "comision"}
                    expected_first = negative if card else not negative
                    if bool(a) != expected_first:
                        raise ValueError("La columna del importe no corresponde al tipo de movimiento")
                    amount = -(a or b) if negative else (a or b)
                    if not str(desc or "").strip():
                        raise ValueError("Falta la descripción")
                    category = str(cat or "Por revisar").strip()
                    categories.add(category)
                    r = {"date": dt.date().isoformat() if isinstance(dt, datetime) else dt.isoformat(),
                        "account": account, "kind": kind, "category": category, "description": str(desc).strip(),
                        "amount_minor": amount, "currency": "MXN", "reference": str(ref).strip() if ref is not None else None,
                        "property": str(prop).strip() if prop is not None else None,
                        "treatment": ledger.treatment(category, kind), "source_file": filename,
                        "source_sheet": name, "source_row": idx, "source_hash": sha}
                    movements.append(r)
                except (ValueError, TypeError) as exc:
                    amount_hint = None
                    try:
                        a_hint, b_hint = ledger.money(incoming or 0), ledger.money(outgoing or 0)
                        if bool(a_hint) != bool(b_hint):
                            amount_hint = abs(a_hint or b_hint)
                    except ValueError:
                        pass
                    errors.append({"sheet": name, "row": idx, "message": str(exc), "draft": {
                        "date": dt.date().isoformat() if isinstance(dt, datetime) else dt.isoformat() if isinstance(dt, date) else "",
                        "account": account, "kind": ledger.normalize(kind), "category": str(cat or "Por revisar").strip(),
                        "description": str(desc or "").strip(), "amount_minor": amount_hint,
                        "reference": str(ref).strip() if ref is not None else None,
                        "property": str(prop).strip() if prop is not None else None}})
        if "Catalogos" in w.sheetnames:
            for row in w["Catalogos"].iter_rows(min_row=2, max_col=3, values_only=True):
                for value in (row[0], row[2]):
                    if isinstance(value, str) and value.strip():
                        categories.add(value.strip())
        if not movements and not errors:
            errors.append({"sheet": "Archivo", "row": 0, "message": "No se encontraron movimientos"})
        return refresh_preview({"filename": filename, "file_hash": sha, "movements": movements, "openings": openings,
            "categories": sorted(categories), "errors": errors, "notes": notes, "corrections": []})
    finally:
        w.close()
